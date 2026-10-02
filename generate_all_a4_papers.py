import json
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_a4_question_and_answer_paper(set_dict, output_path):
    doc = docx.Document()
    
    # Configure A4 Margins (10mm all sides -> 190mm printable width)
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(10)
        section.bottom_margin = Mm(10)
        section.left_margin = Mm(10)
        section.right_margin = Mm(10)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # 1. Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(1)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_brand = title_p.add_run("BEERBOX MIDRAND • LIVE PUB QUIZ NIGHT")
    run_brand.font.size = Pt(15)
    run_brand.font.bold = True
    run_brand.font.color.rgb = RGBColor(180, 83, 9) # Amber 700

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(2)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Official Master Question & Answer Paper • Hosted by DJ JC")
    run_sub.font.size = Pt(9.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    set_title_clean = set_dict.get('setName', f"Set {set_dict.get('setId', 1)}").upper()
    set_p = doc.add_paragraph()
    set_p.paragraph_format.space_before = Pt(0)
    set_p.paragraph_format.space_after = Pt(5)
    set_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_set = set_p.add_run(f"★ {set_title_clean} ★")
    run_set.font.size = Pt(11)
    run_set.font.bold = True
    run_set.font.color.rgb = RGBColor(15, 23, 42)
    if set_dict.get('theme'):
        run_theme = set_p.add_run(f" - {set_dict['theme']}")
        run_theme.font.size = Pt(9)
        run_theme.font.italic = True
        run_theme.font.color.rgb = RGBColor(100, 116, 139)

    # 2. Team Info Table
    team_tbl = doc.add_table(rows=2, cols=4)
    team_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(team_tbl, color="CBD5E1", sz="6")
    
    widths = [Mm(25), Mm(80), Mm(32), Mm(53)]
    for r in team_tbl.rows:
        for idx, width in enumerate(widths):
            r.cells[idx].width = width

    # Row 0: Team Name & Table Number
    cell_t0 = team_tbl.cell(0, 0)
    set_cell_shading(cell_t0, "F8FAFC")
    set_cell_margins(cell_t0, top=70, bottom=70, left=90, right=90)
    p = cell_t0.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r_lbl = p.add_run("Team Name:")
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(9)

    cell_t1 = team_tbl.cell(0, 1)
    set_cell_margins(cell_t1, top=70, bottom=70, left=90, right=90)
    p = cell_t1.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run("________________________________________")

    cell_t2 = team_tbl.cell(0, 2)
    set_cell_shading(cell_t2, "F8FAFC")
    set_cell_margins(cell_t2, top=70, bottom=70, left=90, right=90)
    p = cell_t2.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r_lbl = p.add_run("Table No:")
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(9)

    cell_t3 = team_tbl.cell(0, 3)
    set_cell_margins(cell_t3, top=70, bottom=70, left=90, right=90)
    p = cell_t3.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Table # [ _______ ]")

    # Row 1: Captain & Joker Round
    cell_c0 = team_tbl.cell(1, 0)
    set_cell_shading(cell_c0, "F8FAFC")
    set_cell_margins(cell_c0, top=70, bottom=70, left=90, right=90)
    p = cell_c0.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r_lbl = p.add_run("Captain:")
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(9)

    cell_c1 = team_tbl.cell(1, 1)
    set_cell_margins(cell_c1, top=70, bottom=70, left=90, right=90)
    p = cell_c1.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run("________________________________________")

    cell_c2 = team_tbl.cell(1, 2)
    set_cell_shading(cell_c2, "FEF3C7") # Warm yellow for Joker
    set_cell_margins(cell_c2, top=70, bottom=70, left=90, right=90)
    p = cell_c2.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r_lbl = p.add_run("★ JOKER (x2):")
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(8.5)
    r_lbl.font.color.rgb = RGBColor(180, 83, 9)

    cell_c3 = team_tbl.cell(1, 3)
    set_cell_shading(cell_c3, "FFFBEB")
    set_cell_margins(cell_c3, top=70, bottom=70, left=90, right=90)
    p = cell_c3.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run("[ ] R1  [ ] R2  [ ] R3  [ ] R4  [ ] R5")
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.bold = True

    # Instructions line
    inst_p = doc.add_paragraph()
    inst_p.paragraph_format.space_before = Pt(3)
    inst_p.paragraph_format.space_after = Pt(6)
    inst_run = inst_p.add_run("Master Question & Answer Key: Each question displays the full multiple-choice options with the official verified answer displayed in the Answer column. Strictly No Shazam / Google!")
    inst_run.font.italic = True
    inst_run.font.size = Pt(8)
    inst_run.font.color.rgb = RGBColor(100, 116, 139)

    # 3. Render Each Round
    rounds = set_dict.get('rounds', [])
    for r_idx, r in enumerate(rounds):
        is_tiebreaker = (r.get("roundId") == 6)
        
        # Round Header Table / Banner
        banner_tbl = doc.add_table(rows=1, cols=2)
        banner_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        banner_cell_left = banner_tbl.cell(0, 0)
        banner_cell_right = banner_tbl.cell(0, 1)
        banner_cell_left.width = Mm(145)
        banner_cell_right.width = Mm(45)
        
        bg_color = "7E22CE" if is_tiebreaker else "1E293B" # Purple for tiebreaker, Slate for rounds
        set_cell_shading(banner_cell_left, bg_color)
        set_cell_shading(banner_cell_right, bg_color)
        set_cell_margins(banner_cell_left, top=80, bottom=80, left=100, right=100)
        set_cell_margins(banner_cell_right, top=80, bottom=80, left=100, right=100)
        
        p_left = banner_cell_left.paragraphs[0]
        p_left.paragraph_format.space_after = Pt(0)
        r_title = p_left.add_run(f"{r['title'].upper()}")
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(254, 243, 199) # Warm cream
        
        if r.get('theme'):
            p_sub = banner_cell_left.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(1)
            p_sub.paragraph_format.space_after = Pt(0)
            r_theme = p_sub.add_run(r['theme'])
            r_theme.font.size = Pt(8)
            r_theme.font.color.rgb = RGBColor(203, 213, 225)
        
        p_right = banner_cell_right.paragraphs[0]
        p_right.paragraph_format.space_after = Pt(0)
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_badge = p_right.add_run(r.get('badge', ''))
        r_badge.font.bold = True
        r_badge.font.size = Pt(8.5)
        r_badge.font.color.rgb = RGBColor(251, 191, 36) # Amber gold

        # Question Table: Col 0 (#), Col 1 (Question & Options), Col 2 (Answer), Col 3 (Host Notes)
        q_count = len(r.get('questions', []))
        q_tbl = doc.add_table(rows=q_count + 1, cols=4)
        q_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(q_tbl, color="E2E8F0", sz="4")
        
        col_widths = [Mm(9), Mm(98), Mm(41), Mm(42)]
        for row in q_tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = w

        # Header Row
        hdr_row = q_tbl.rows[0]
        headers = ["#", "Question & Multiple-Choice Options", "Official Answer", "Host Notes / Trivia"]
        aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]
        
        for idx, text in enumerate(headers):
            cell = hdr_row.cells[idx]
            set_cell_shading(cell, "F1F5F9")
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = aligns[idx]
            run = p.add_run(text)
            run.font.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(51, 65, 85)

        # Question Rows
        for q_idx, q in enumerate(r.get('questions', [])):
            row = q_tbl.rows[q_idx + 1]
            
            # Col 0: Q Number
            c0 = row.cells[0]
            set_cell_margins(c0, top=45, bottom=45, left=50, right=50)
            c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run0 = p0.add_run(f"Q{q_idx + 1}")
            run0.font.bold = True
            run0.font.size = Pt(8.5)
            run0.font.color.rgb = RGBColor(100, 116, 139)
            
            # Col 1: Question Text & Options
            c1 = row.cells[1]
            set_cell_margins(c1, top=45, bottom=45, left=70, right=70)
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(1)
            run_q = p1.add_run(q['q'])
            run_q.font.bold = True
            run_q.font.size = Pt(8.5)
            run_q.font.color.rgb = RGBColor(15, 23, 42)
            
            # Multiple Choice Options
            if 'options' in q and q['options']:
                p_opts = c1.add_paragraph()
                p_opts.paragraph_format.space_before = Pt(0)
                p_opts.paragraph_format.space_after = Pt(0)
                opts_text = "   ".join([f"○ {opt}" for opt in q['options']])
                run_opts = p_opts.add_run(opts_text)
                run_opts.font.size = Pt(7.5)
                run_opts.font.color.rgb = RGBColor(71, 85, 105)

            # Col 2: Official Answer (Highlight in Emerald Soft Green)
            c2 = row.cells[2]
            set_cell_shading(c2, "ECFDF5")
            set_cell_margins(c2, top=45, bottom=45, left=70, right=70)
            c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            run_ans = p2.add_run(f"✓ {q.get('a', '')}")
            run_ans.font.bold = True
            run_ans.font.size = Pt(8.5)
            run_ans.font.color.rgb = RGBColor(4, 120, 87) # Emerald 700
            
            # Col 3: Host Notes / Trivia
            c3 = row.cells[3]
            set_cell_margins(c3, top=45, bottom=45, left=60, right=60)
            c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p3 = c3.paragraphs[0]
            p3.paragraph_format.space_after = Pt(0)
            run_note = p3.add_run(q.get('notes', ''))
            run_note.font.italic = True
            run_note.font.size = Pt(7.5)
            run_note.font.color.rgb = RGBColor(100, 116, 139)

        # Round Subtotal Bar
        subtotal_tbl = doc.add_table(rows=1, cols=3)
        subtotal_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(subtotal_tbl, color="CBD5E1", sz="6")
        
        st_widths = [Mm(62), Mm(66), Mm(62)]
        for idx, w in enumerate(st_widths):
            subtotal_tbl.cell(0, idx).width = w

        c_st0 = subtotal_tbl.cell(0, 0)
        set_cell_shading(c_st0, "FEF3C7") # Light amber
        set_cell_margins(c_st0, top=60, bottom=60, left=90, right=90)
        p = c_st0.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(f"Round {r.get('roundId', r_idx+1)} Base Score:  [ ____ / {q_count} ]")
        r1.font.bold = True
        r1.font.size = Pt(8.5)

        c_st1 = subtotal_tbl.cell(0, 1)
        set_cell_shading(c_st1, "FFFBEB")
        set_cell_margins(c_st1, top=60, bottom=60, left=90, right=90)
        p = c_st1.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if not is_tiebreaker:
            r2 = p.add_run("Joker Active?  [ ] YES (x2)   [ ] NO")
        else:
            r2 = p.add_run("Tie-Breaker: Closest Estimate")
        r2.font.bold = True
        r2.font.size = Pt(8.5)

        c_st2 = subtotal_tbl.cell(0, 2)
        set_cell_shading(c_st2, "FDE68A") # Darker warm amber
        set_cell_margins(c_st2, top=60, bottom=60, left=90, right=90)
        p = c_st2.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r3 = p.add_run(f"Round {r.get('roundId', r_idx+1)} Official Total:  [ _______ ]")
        r3.font.bold = True
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = RGBColor(146, 64, 14)

        # Spacing after round
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_before = Pt(0)
        sp_p.paragraph_format.space_after = Pt(5)

        # Page breaks after Round 2 and Round 4
        if r.get('roundId') == 2 or r.get('roundId') == 4:
            doc.add_page_break()

    # 4. Master Grand Total Calculation Scorecard
    gt_banner = doc.add_table(rows=1, cols=1)
    gt_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = gt_banner.cell(0, 0)
    c.width = Mm(190)
    set_cell_shading(c, "0F172A") # Midnight
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_gt = p.add_run("★ OFFICIAL GRAND TOTAL SCORECARD (END OF GAME) ★")
    r_gt.font.bold = True
    r_gt.font.size = Pt(10.5)
    r_gt.font.color.rgb = RGBColor(251, 191, 36)

    gt_tbl = doc.add_table(rows=3, cols=6)
    gt_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(gt_tbl, color="CBD5E1", sz="6")
    
    gt_col_w = Mm(31.6)
    for row in gt_tbl.rows:
        for idx in range(6):
            row.cells[idx].width = gt_col_w

    # Row 0: Labels
    round_labels = ["Round 1", "Round 2", "Round 3", "Round 4", "Round 5", "Joker Bonus"]
    for idx, lbl in enumerate(round_labels):
        cell = gt_tbl.cell(0, idx)
        set_cell_shading(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lbl)
        r.font.bold = True
        r.font.size = Pt(8.5)

    # Row 1: Score Inputs
    for idx in range(5):
        cell = gt_tbl.cell(1, idx)
        set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("[   / 10 ]")
        r.font.bold = True
        r.font.size = Pt(9.5)

    cell_jb = gt_tbl.cell(1, 5)
    set_cell_shading(cell_jb, "FEF3C7")
    set_cell_margins(cell_jb, top=70, bottom=70, left=70, right=70)
    p = cell_jb.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("[ + ____ ]")
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(180, 83, 9)

    # Row 2: Final Grand Total spanning across
    row2 = gt_tbl.rows[2]
    c_final = row2.cells[0]
    c_final.merge(row2.cells[3])
    set_cell_shading(c_final, "FEF3C7")
    set_cell_margins(c_final, top=80, bottom=80, left=100, right=100)
    p = c_final.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("FINAL TOURNAMENT SCORE:   [ ____________ / 50 ]")
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(146, 64, 14)

    c_rank = row2.cells[1]
    c_rank.merge(row2.cells[2])
    set_cell_shading(c_rank, "FDE68A")
    set_cell_margins(c_rank, top=80, bottom=80, left=100, right=100)
    p = c_rank.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FINAL RANK:  # [ ______ ]")
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(146, 64, 14)

    # Signoff line
    so_p = doc.add_paragraph()
    so_p.paragraph_format.space_before = Pt(6)
    so_p.paragraph_format.space_after = Pt(0)
    so_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_so = so_p.add_run("Official Marker Signature: _______________________      DJ JC Verified: [  ✓  ]")
    r_so.font.size = Pt(8)
    r_so.font.italic = True
    r_so.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(str(output_path))
    print(f"Generated A4 Paper with Answers: {output_path.name}")

def main():
    repo_dir = Path("/home/jakes/repos/quiznight")
    json_path = repo_dir / "quizSets.json"
    
    if not json_path.exists():
        subprocess.run(["node", "-e", "const fs=require('fs'); eval(fs.readFileSync('quizSets.js','utf8')); fs.writeFileSync('quizSets.json', JSON.stringify(quizSets, null, 2));"], cwd=str(repo_dir), check=True)
    
    with open(json_path, "r", encoding="utf-8") as f:
        quiz_sets = json.load(f)
    
    print(f"Found {len(quiz_sets)} quiz sets to process.")
    
    for s in quiz_sets:
        set_id = s.get('setId')
        out_file = repo_dir / f"Beerbox_Pub_Quiz_Set_{set_id}_A4_With_Answers.docx"
        build_a4_question_and_answer_paper(s, out_file)
        
    # Overwrite default Beerbox_Pub_Quiz_Team_Question_Paper_A4.docx with Set 1 with answers
    default_paper = repo_dir / "Beerbox_Pub_Quiz_Team_Question_Paper_A4.docx"
    build_a4_question_and_answer_paper(quiz_sets[0], default_paper)
    print(f"Updated default paper: {default_paper.name}")

if __name__ == "__main__":
    main()
